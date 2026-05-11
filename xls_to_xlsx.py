#!/usr/bin/env python3
"""
批量将 .xls 文件转换为 .xlsx 格式
支持真正的 Excel 二进制格式和伪装成 .xls 的 TSV/CSV 文本文件
转换后的文件输出到源目录下的 Convert 文件夹
用法:
    Windows : py xls_to_xlsx.py [目录路径]   或双击 / 拖拽 run.bat
    macOS   : python3 xls_to_xlsx.py [目录路径]
"""

import sys
import os
import glob
import csv
import subprocess

# 让脚本在 Windows 中文 cmd（GBK 代码页）下也能正常打印 emoji / 中文
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


# ==============================================================
# 环境自检
# ==============================================================

def _print_upgrade_hint():
    if sys.platform == "win32":
        print("\n⚠️  Python 版本过低，请到 https://www.python.org/downloads/ 下载并安装 3.11+")
        print("    安装时务必勾选 “Add python.exe to PATH”")
        print("    安装后重新打开命令行或双击 run.bat 即可")
    elif sys.platform == "darwin":
        print("\n⚠️  Python 版本过低，建议升级到 3.11：")
        print("    brew install python@3.11")
        print('    echo \'alias python3="/opt/homebrew/opt/python@3.11/bin/python3.11"\' >> ~/.zshrc')
        print('    echo \'alias pip3="/opt/homebrew/opt/python@3.11/bin/pip3.11"\' >> ~/.zshrc')
        print("    source ~/.zshrc")
        print("\n    然后重新运行此脚本。")
    else:
        print("\n⚠️  Python 版本过低，请升级到 Python 3.9 或更高版本后重试。")


def check_python_version():
    """检查 Python 版本"""
    v = sys.version_info
    print(f"[检查] Python 版本: {v.major}.{v.minor}.{v.micro}")
    if v.major < 3 or (v.major == 3 and v.minor < 9):
        _print_upgrade_hint()
        sys.exit(1)


def check_and_install_deps():
    """检查并自动安装依赖库"""
    missing = []
    for pkg in ["xlrd", "openpyxl"]:
        try:
            __import__(pkg)
        except ImportError:
            missing.append(pkg)

    if not missing:
        return True

    print(f"\n[检查] 缺少依赖: {', '.join(missing)}")
    print(f"       正在自动安装...")

    try:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", "--break-system-packages", *missing],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
        )
    except subprocess.CalledProcessError:
        try:
            subprocess.check_call(
                [sys.executable, "-m", "pip", "install", *missing],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE,
            )
        except subprocess.CalledProcessError:
            print("\n❌ 自动安装失败，请手动执行：")
            if sys.platform == "win32":
                print(f"    py -m pip install {' '.join(missing)}")
            else:
                print(f"    pip3 install {' '.join(missing)}")
                print("    或")
                print(f"    python3 -m pip install {' '.join(missing)}")
            sys.exit(1)

    for pkg in missing:
        try:
            __import__(pkg)
        except ImportError:
            print(f"\n❌ {pkg} 安装后仍无法导入，请手动安装后重试。")
            sys.exit(1)

    print(f"       ✅ 已自动安装: {', '.join(missing)}")
    return True


# 运行自检
check_python_version()
check_and_install_deps()

import xlrd
import openpyxl

# ==============================================================
# 转换逻辑
# ==============================================================

# 文本编码探测顺序：BOM 优先，再 UTF-8，再中文 Windows 常见编码
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gbk", "gb18030")


def is_real_xls(filepath):
    with open(filepath, "rb") as f:
        header = f.read(8)
    return header[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def read_text_with_fallback(filepath):
    """按 TEXT_ENCODINGS 顺序尝试解码；都失败则用 utf-8 + replace 兜底。返回 (text, encoding)。"""
    with open(filepath, "rb") as f:
        raw = f.read()
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8(replace)"


def detect_delimiter(sample):
    tab_count = sample.count("\t")
    comma_count = sample.count(",")
    return "\t" if tab_count >= comma_count else ","


def convert_text_to_xlsx(filepath, output_dir):
    basename = os.path.splitext(os.path.basename(filepath))[0]
    xlsx_path = os.path.join(output_dir, f"{basename}.xlsx")

    text, encoding = read_text_with_fallback(filepath)
    delimiter = detect_delimiter(text[:4096])
    fmt = "TSV" if delimiter == "\t" else "CSV"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    for row_idx, row in enumerate(reader, 1):
        for col_idx, value in enumerate(row, 1):
            try:
                value = int(value)
            except ValueError:
                try:
                    value = float(value)
                except ValueError:
                    pass
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(xlsx_path)
    return xlsx_path, fmt, encoding


def convert_xls_to_xlsx(xls_path, output_dir):
    basename = os.path.splitext(os.path.basename(xls_path))[0]
    xlsx_path = os.path.join(output_dir, f"{basename}.xlsx")

    xls_book = xlrd.open_workbook(xls_path, formatting_info=False)
    xlsx_book = openpyxl.Workbook()
    xlsx_book.remove(xlsx_book.active)

    for sheet_name in xls_book.sheet_names():
        xls_sheet = xls_book.sheet_by_name(sheet_name)
        xlsx_sheet = xlsx_book.create_sheet(title=sheet_name)

        for row in range(xls_sheet.nrows):
            for col in range(xls_sheet.ncols):
                cell = xls_sheet.cell(row, col)
                value = cell.value
                xlsx_cell = xlsx_sheet.cell(row=row + 1, column=col + 1)

                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        dt_tuple = xlrd.xldate_as_tuple(value, xls_book.datemode)
                        if dt_tuple[0] == 0 and dt_tuple[1] == 0 and dt_tuple[2] == 0:
                            # 纯时间 → 写成文本 "HH:MM:SS"
                            xlsx_cell.value = f"{dt_tuple[3]:02d}:{dt_tuple[4]:02d}:{dt_tuple[5]:02d}"
                        else:
                            # 有日期部分 → 写成文本 "YYYY-MM-DD"
                            xlsx_cell.value = f"{dt_tuple[0]:04d}-{dt_tuple[1]:02d}-{dt_tuple[2]:02d}"
                    except Exception:
                        xlsx_cell.value = value
                else:
                    xlsx_cell.value = value

        # 自动调整列宽
        for col_cells in xlsx_sheet.columns:
            max_len = 0
            col_letter = col_cells[0].column_letter
            for c in col_cells:
                if c.value is not None:
                    max_len = max(max_len, len(str(c.value)))
            xlsx_sheet.column_dimensions[col_letter].width = max_len + 2

    xlsx_book.save(xlsx_path)
    return xlsx_path


def main():
    if len(sys.argv) > 1:
        target_dir = sys.argv[1]
    else:
        target_dir = os.path.dirname(os.path.abspath(__file__))

    target_dir = os.path.abspath(target_dir)

    # 创建 Convert 输出目录
    output_dir = os.path.join(target_dir, "Convert")
    os.makedirs(output_dir, exist_ok=True)

    xls_files = glob.glob(os.path.join(target_dir, "*.xls"))
    xls_files = [f for f in xls_files if f.lower().endswith(".xls")]

    if not xls_files:
        print(f"\n在 {target_dir} 下未找到 .xls 文件")
        return

    print(f"\n找到 {len(xls_files)} 个 .xls 文件，开始转换...")
    print(f"输出目录: {output_dir}\n")

    success, fail = 0, 0
    for xls_path in sorted(xls_files):
        name = os.path.basename(xls_path)
        try:
            if is_real_xls(xls_path):
                out = convert_xls_to_xlsx(xls_path, output_dir)
                print(f"  ✅ {name} → {os.path.basename(out)}  [Excel 二进制]")
            else:
                out, fmt, encoding = convert_text_to_xlsx(xls_path, output_dir)
                print(f"  ✅ {name} → {os.path.basename(out)}  [实际为 {fmt} 文本, 编码 {encoding}]")
            success += 1
        except Exception as e:
            print(f"  ❌ {name} 失败: {e}")
            fail += 1

    print(f"\n完成: {success} 成功, {fail} 失败")
    if success > 0:
        print(f"文件已保存到: {output_dir}")


if __name__ == "__main__":
    main()
